#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Aggregate the LOI study's results into:
  * results/loi_aggregate_metrics.csv  — mean [95% CI] for each LOI config + reference rows
  * figures/loi_delta_auc.{png,pdf}    — horizontal bar chart of ΔAUC vs Full and ΔAUC vs context_only
  * figures/loi_metric_grid.{png,pdf}  — 4-panel AUC/PR-AUC/F1/MCC for the 4 LOI configs + 3 references
  * report/loi_report.{md,docx}        — narrative report with table + figures

Reference rows (Full CARCM, sequence_only, context_only) are reused from the
parent study's `../results/{full,sequence_only,context_only}/` (no retraining).
"""

import argparse
import os
import sys

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib as mpl

THIS_DIR = os.path.dirname(os.path.abspath(__file__))
PARENT_DIR = os.path.dirname(THIS_DIR)
sys.path.insert(0, THIS_DIR)
sys.path.insert(0, PARENT_DIR)

from ablation_utils import delong_roc_test, seed_mean_ci, metric_pack  # noqa: E402
from aggregate_ablation import best_ensemble  # noqa: E402  reuse parent's
from loi_configs import LOI_CONFIGS  # noqa: E402

mpl.rcParams.update({
    "font.size": 11, "font.family": "sans-serif",
    "axes.labelsize": 12, "axes.titlesize": 12,
    "xtick.labelsize": 10, "ytick.labelsize": 10,
    "legend.fontsize": 10,
    "figure.dpi": 400, "savefig.dpi": 400, "savefig.bbox": "tight",
    "axes.grid": True, "grid.alpha": 0.3,
})


REFERENCE_ROWS = [
    {"id": 1, "name": "full", "label": "Full CARCM (all 4 features + context)",
     "_results_root": os.path.join(PARENT_DIR, "results")},
    {"id": 2, "name": "sequence_only", "label": "Sequence-only (no context)",
     "_results_root": os.path.join(PARENT_DIR, "results")},
    {"id": 7, "name": "context_only", "label": "Context-only (no sequence features)",
     "_results_root": os.path.join(PARENT_DIR, "results")},
]


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("--root", type=str, default=THIS_DIR + "/")
    return p.parse_args()


def load_per_seed(results_root, name):
    p = os.path.join(results_root, name, "per_seed_metrics.csv")
    if not os.path.exists(p):
        return None
    df = pd.read_csv(p)
    df["config"] = name
    return df


def load_predictions(results_root, name, split="test"):
    sub = "predictions" if split == "test" else "predictions_dev"
    pdir = os.path.join(results_root, name, sub)
    if not os.path.isdir(pdir):
        return None
    files = sorted(f for f in os.listdir(pdir) if f.endswith(".npy"))
    if not files:
        return None
    return np.stack([np.load(os.path.join(pdir, f)) for f in files], axis=0)


def significance_marker(p):
    if p is None or not np.isfinite(p):
        return "n.s."
    if p < 1e-4: return "****"
    if p < 1e-3: return "***"
    if p < 1e-2: return "**"
    if p < 5e-2: return "*"
    return "n.s."


def fmt_p(p):
    if p is None or not np.isfinite(p):
        return "n/a"
    if p == 0.0:
        return "<1e-300"
    if p < 1e-4:
        return f"{p:.2e}"
    return f"{p:.4f}"


def aggregate(args):
    root = args.root
    results_root = os.path.join(root, "results")
    fig_dir = os.path.join(root, "figures")
    rep_dir = os.path.join(root, "report")
    os.makedirs(fig_dir, exist_ok=True)
    os.makedirs(rep_dir, exist_ok=True)

    # ---------- assemble configs to aggregate ----------
    all_cfgs = []
    for r in REFERENCE_ROWS:
        all_cfgs.append({**r, "_is_reference": True})
    for c in LOI_CONFIGS:
        all_cfgs.append({**c, "_is_reference": False,
                         "_results_root": results_root})

    # ---------- load per-seed metrics ----------
    per_config_seed = {}
    long_rows = []
    for cfg in all_cfgs:
        df = load_per_seed(cfg["_results_root"], cfg["name"])
        if df is None:
            print(f"WARN: {cfg['name']} has no per_seed_metrics.csv yet", flush=True)
            continue
        df = df.copy()
        df["id"] = cfg["id"]
        df["label"] = cfg["label"]
        per_config_seed[cfg["name"]] = df
        long_rows.append(df)
    if not long_rows:
        raise RuntimeError("No per-seed results found.")
    long_df = pd.concat(long_rows, ignore_index=True)
    long_df.to_csv(os.path.join(results_root, "loi_per_seed_long.csv"), index=False)

    # ---------- load test labels (canonical, from parent's full) ----------
    test_labels = np.load(os.path.join(REFERENCE_ROWS[0]["_results_root"],
                                       "full", "test_labels.npy"))

    # ---------- build v7b 5-method ensemble per config ----------
    ensemble = {}
    method_used = {}
    method_aucs = {}
    ensemble_dev = {}          # dev-side prediction of each config's chosen ensemble
    ensemble_dev_labels = {}
    for cfg in all_cfgs:
        test_preds = load_predictions(cfg["_results_root"], cfg["name"], split="test")
        dev_preds = load_predictions(cfg["_results_root"], cfg["name"], split="dev")
        if test_preds is None:
            continue
        if dev_preds is None:
            ensemble[cfg["name"]] = test_preds.mean(axis=0)
            method_used[cfg["name"]] = "simple_avg"
            method_aucs[cfg["name"]] = {}
            continue
        dev_labels_path = os.path.join(cfg["_results_root"], cfg["name"], "dev_labels.npy")
        dev_labels = np.load(dev_labels_path)
        best_name, best_preds, _, mauc, best_dev_preds = best_ensemble(
            test_preds, dev_preds, dev_labels, test_labels
        )
        ensemble[cfg["name"]] = best_preds
        ensemble_dev[cfg["name"]] = best_dev_preds
        ensemble_dev_labels[cfg["name"]] = dev_labels
        method_used[cfg["name"]] = best_name
        method_aucs[cfg["name"]] = mauc

    pd.DataFrame([
        {"name": n, "best_method": method_used.get(n, ""),
         **{f"{m}_{k}": v for m, d in method_aucs.get(n, {}).items()
            for k, v in d.items()}}
        for n in ensemble
    ]).to_csv(os.path.join(results_root, "loi_ensemble_selection.csv"), index=False)

    # ---------- assemble row-level table ----------
    rows = []
    for cfg in all_cfgs:
        if cfg["name"] not in per_config_seed:
            continue
        d = per_config_seed[cfg["name"]]
        row = {
            "id": cfg["id"], "name": cfg["name"], "label": cfg["label"],
            "n_seeds": len(d),
            "is_reference": cfg.get("_is_reference", False),
        }
        for metric in ["auc", "pr_auc", "f1", "mcc"]:
            m, lo, hi = seed_mean_ci(d[metric].values, ci=95)
            row[f"{metric}_mean"] = m
            row[f"{metric}_lo"] = lo
            row[f"{metric}_hi"] = hi
            row[f"{metric}_str"] = f"{m:.4f} [{lo:.4f}, {hi:.4f}]"
        if cfg["name"] in ensemble:
            ens_m = metric_pack(
                test_labels, ensemble[cfg["name"]],
                dev_labels=ensemble_dev_labels.get(cfg["name"]),
                dev_preds=ensemble_dev.get(cfg["name"]))
            row["auc_ensemble"] = ens_m["auc"]
            row["pr_auc_ensemble"] = ens_m["pr_auc"]
            row["f1_ensemble"] = ens_m["f1"]
            row["mcc_ensemble"] = ens_m["mcc"]
            row["threshold_ensemble"] = ens_m["threshold"]
        rows.append(row)
    agg = pd.DataFrame(rows)

    # ---------- ΔAUC vs Full (paired across seeds) ----------
    if "full" not in per_config_seed:
        raise RuntimeError("Full CARCM per-seed metrics missing.")
    full_aucs = per_config_seed["full"]["auc"].values
    delta_full_mean, delta_full_lo, delta_full_hi = {}, {}, {}
    for cfg in all_cfgs:
        if cfg["name"] not in per_config_seed:
            continue
        this_aucs = per_config_seed[cfg["name"]]["auc"].values
        n = min(len(this_aucs), len(full_aucs))
        diff = this_aucs[:n] - full_aucs[:n]
        m, lo, hi = seed_mean_ci(diff, ci=95)
        delta_full_mean[cfg["name"]] = m
        delta_full_lo[cfg["name"]] = lo
        delta_full_hi[cfg["name"]] = hi
    agg["delta_auc_vs_full_mean"] = agg["name"].map(delta_full_mean)
    agg["delta_auc_vs_full_lo"] = agg["name"].map(delta_full_lo)
    agg["delta_auc_vs_full_hi"] = agg["name"].map(delta_full_hi)
    agg["delta_auc_vs_full_str"] = agg.apply(
        lambda r: f"{r['delta_auc_vs_full_mean']:+.4f} "
                  f"[{r['delta_auc_vs_full_lo']:+.4f}, {r['delta_auc_vs_full_hi']:+.4f}]",
        axis=1,
    )

    # ---------- ΔAUC vs context_only (the "what does this feature ADD to context?") ----------
    if "context_only" in per_config_seed:
        ctxo_aucs = per_config_seed["context_only"]["auc"].values
        d_ctxo_mean, d_ctxo_lo, d_ctxo_hi = {}, {}, {}
        for cfg in all_cfgs:
            if cfg["name"] not in per_config_seed:
                continue
            this_aucs = per_config_seed[cfg["name"]]["auc"].values
            n = min(len(this_aucs), len(ctxo_aucs))
            diff = this_aucs[:n] - ctxo_aucs[:n]
            m, lo, hi = seed_mean_ci(diff, ci=95)
            d_ctxo_mean[cfg["name"]] = m
            d_ctxo_lo[cfg["name"]] = lo
            d_ctxo_hi[cfg["name"]] = hi
        agg["delta_auc_vs_context_only_mean"] = agg["name"].map(d_ctxo_mean)
        agg["delta_auc_vs_context_only_lo"] = agg["name"].map(d_ctxo_lo)
        agg["delta_auc_vs_context_only_hi"] = agg["name"].map(d_ctxo_hi)
        agg["delta_auc_vs_context_only_str"] = agg.apply(
            lambda r: f"{r['delta_auc_vs_context_only_mean']:+.4f} "
                      f"[{r['delta_auc_vs_context_only_lo']:+.4f}, "
                      f"{r['delta_auc_vs_context_only_hi']:+.4f}]",
            axis=1,
        )

    # ---------- DeLong p-value vs Full and vs context_only on ensemble preds ----------
    delong_rows = []
    p_full_map, z_full_map = {}, {}
    p_ctxo_map, z_ctxo_map = {}, {}
    full_preds = ensemble.get("full")
    ctxo_preds = ensemble.get("context_only")
    for cfg in all_cfgs:
        n = cfg["name"]
        if n not in ensemble:
            continue
        a, b, zF, pF = delong_roc_test(test_labels, full_preds, ensemble[n]) if n != "full" \
            else (float("nan"), float("nan"), float("nan"), float("nan"))
        if ctxo_preds is not None and n != "context_only":
            a2, b2, zC, pC = delong_roc_test(test_labels, ctxo_preds, ensemble[n])
        else:
            zC, pC = float("nan"), float("nan")
        delong_rows.append({"name": n, "auc_full": a, "auc_other": b,
                            "z_vs_full": zF, "p_vs_full": pF,
                            "z_vs_context_only": zC, "p_vs_context_only": pC})
        p_full_map[n] = pF
        z_full_map[n] = zF
        p_ctxo_map[n] = pC
        z_ctxo_map[n] = zC
    pd.DataFrame(delong_rows).to_csv(
        os.path.join(results_root, "loi_delong.csv"), index=False)
    agg["delong_p_vs_full"] = agg["name"].map(p_full_map)
    agg["delong_p_vs_context_only"] = agg["name"].map(p_ctxo_map)
    agg["delong_marker_vs_full"] = agg["delong_p_vs_full"].apply(significance_marker)
    agg["delong_marker_vs_context_only"] = agg["delong_p_vs_context_only"].apply(significance_marker)

    agg = agg.sort_values("id").reset_index(drop=True)
    agg.to_csv(os.path.join(results_root, "loi_aggregate_metrics.csv"), index=False)

    # ---------- figures ----------
    make_loi_main_figure(agg, fig_dir)
    make_loi_metric_grid(agg, fig_dir)

    # ---------- report ----------
    write_report(agg, root, n_seeds=int(long_df.groupby("config")["seed"].nunique().max()),
                 n_test=len(test_labels))
    print("LOI AGGREGATION COMPLETE", flush=True)
    print(f"  csv:  {os.path.join(results_root, 'loi_aggregate_metrics.csv')}", flush=True)
    print(f"  fig:  {os.path.join(fig_dir, 'loi_delta_auc.png')}", flush=True)
    print(f"  rpt:  {os.path.join(rep_dir, 'loi_report.md')}", flush=True)


def make_loi_main_figure(agg, fig_dir):
    """Two-panel: ΔAUC vs Full (left, what each LOI loses) and
    ΔAUC vs context_only (right, what each LOI adds on top of context)."""
    loi_names = ["only_binned", "only_positional", "only_count", "only_cooc"]
    plot_df = agg[agg["name"].isin(loi_names)].copy()
    if plot_df.empty:
        return
    plot_df = plot_df.sort_values("delta_auc_vs_full_mean", ascending=True).reset_index(drop=True)

    fig, axes = plt.subplots(1, 2, figsize=(13, 4))
    # Left panel: vs Full (negative — how much each LOI loses vs Full)
    ax = axes[0]
    means = plot_df["delta_auc_vs_full_mean"].values
    los = plot_df["delta_auc_vs_full_lo"].values
    his = plot_df["delta_auc_vs_full_hi"].values
    err_lo = means - los; err_hi = his - means
    y = np.arange(len(plot_df))
    ax.barh(y, means, xerr=[err_lo, err_hi], color="#c66",
            edgecolor="black", linewidth=0.7,
            error_kw=dict(ecolor="black", capsize=3, lw=0.8))
    ax.axvline(0.0, color="black", lw=0.8, alpha=0.6)
    ax.set_yticks(y); ax.set_yticklabels(plot_df["label"].values)
    ax.set_xlabel("ΔAUC vs. Full CARCM (mean ± 95% CI across 10 seeds)")
    ax.set_title("How much each single feature LOSES vs Full")
    for i, marker in enumerate(plot_df["delong_marker_vs_full"].values):
        ax.text(min(means[i], los[i]) - 0.005, i, marker, va="center", ha="right", fontsize=10)

    # Right panel: vs context_only (positive — how much each LOI adds on top of context)
    if "delta_auc_vs_context_only_mean" in plot_df.columns:
        ax = axes[1]
        plot_df_r = plot_df.sort_values("delta_auc_vs_context_only_mean", ascending=True).reset_index(drop=True)
        means = plot_df_r["delta_auc_vs_context_only_mean"].values
        los = plot_df_r["delta_auc_vs_context_only_lo"].values
        his = plot_df_r["delta_auc_vs_context_only_hi"].values
        err_lo = means - los; err_hi = his - means
        y = np.arange(len(plot_df_r))
        ax.barh(y, means, xerr=[err_lo, err_hi], color="#5b8def",
                edgecolor="black", linewidth=0.7,
                error_kw=dict(ecolor="black", capsize=3, lw=0.8))
        ax.axvline(0.0, color="black", lw=0.8, alpha=0.6)
        ax.set_yticks(y); ax.set_yticklabels(plot_df_r["label"].values)
        ax.set_xlabel("ΔAUC vs. context_only (mean ± 95% CI across 10 seeds)")
        ax.set_title("How much each single feature ADDS to context")
        for i, marker in enumerate(plot_df_r["delong_marker_vs_context_only"].values):
            ax.text(max(means[i], his[i]) + 0.005, i, marker, va="center", ha="left", fontsize=10)
    plt.tight_layout()
    for ext in ("png", "pdf"):
        fig.savefig(os.path.join(fig_dir, f"loi_delta_auc.{ext}"))
    plt.close(fig)


def make_loi_metric_grid(agg, fig_dir):
    keep = ["full", "only_binned", "only_positional", "only_count", "only_cooc",
            "sequence_only", "context_only"]
    plot_df = agg[agg["name"].isin(keep)].copy()
    if plot_df.empty:
        return
    order = {n: i for i, n in enumerate(keep)}
    plot_df["__order"] = plot_df["name"].map(order)
    plot_df = plot_df.sort_values("__order").reset_index(drop=True)

    fig, axes = plt.subplots(2, 2, figsize=(11, 6.5))
    metrics = [("auc", "Test AUC"), ("pr_auc", "Test PR-AUC"),
               ("f1", "Test F1"), ("mcc", "Test MCC")]
    y = np.arange(len(plot_df))[::-1]
    for ax, (col, title) in zip(axes.flat, metrics):
        m = plot_df[f"{col}_mean"].values
        lo = plot_df[f"{col}_lo"].values
        hi = plot_df[f"{col}_hi"].values
        err_lo = m - lo; err_hi = hi - m
        ax.barh(y, m, xerr=[err_lo, err_hi], color="#3a6ea5",
                edgecolor="black", linewidth=0.6,
                error_kw=dict(ecolor="black", capsize=3, lw=0.8))
        ax.set_yticks(y); ax.set_yticklabels(plot_df["label"].values, fontsize=8)
        ax.set_xlabel(title); ax.set_title(title)
    plt.tight_layout()
    for ext in ("png", "pdf"):
        fig.savefig(os.path.join(fig_dir, f"loi_metric_grid.{ext}"))
    plt.close(fig)


def write_report(agg, root, n_seeds, n_test):
    rep_dir = os.path.join(root, "report")
    fig_dir = os.path.join(root, "figures")
    os.makedirs(rep_dir, exist_ok=True)

    keep = ["full", "only_binned", "only_positional", "only_count", "only_cooc",
            "sequence_only", "context_only"]
    view = agg[agg["name"].isin(keep)].copy()
    order = {n: i for i, n in enumerate(keep)}
    view["__order"] = view["name"].map(order)
    view = view.sort_values("__order").reset_index(drop=True)

    lines = []
    lines.append("# LOI Feature Ablation — Report\n")
    lines.append("## Executive summary\n")
    lines.append(
        f"Four leave-one-IN configurations: each retains a single sequence-side "
        f"feature block (Binned Motif Counts, Positional Features, Global Motif "
        f"Counts, or Co-occurrence) and the full (Variety, Tissue, Stage) "
        f"context branch. References (Full CARCM, sequence_only, context_only) "
        f"are reused from the parent context-ablation study without retraining. "
        f"Each configuration trained with {n_seeds} seeds, 150 epochs, "
        f"patience 25 — identical to the v7c regime. Metrics are mean [95% CI] "
        f"across seeds; ensemble AUCs come from the v7b 5-method best-of "
        f"selection. Test set: n = {n_test}.\n"
    )

    lines.append("## Table — Single-feature + context (LOI)\n")
    lines.append(
        "| # | Configuration | Ensemble AUC | AUC mean [95% CI] | "
        "ΔAUC vs Full | ΔAUC vs context_only | DeLong p vs Full | "
        "DeLong p vs context_only |"
    )
    lines.append("|---|---|---|---|---|---|---|---|")
    for _, r in view.iterrows():
        e = r.get("auc_ensemble", float("nan"))
        d_full = r.get("delta_auc_vs_full_str", "")
        d_ctxo = r.get("delta_auc_vs_context_only_str", "")
        p_full = r.get("delong_p_vs_full", float("nan"))
        p_ctxo = r.get("delong_p_vs_context_only", float("nan"))
        m_full = r.get("delong_marker_vs_full", "")
        m_ctxo = r.get("delong_marker_vs_context_only", "")
        lines.append(
            f"| {int(r['id'])} | {r['label']} | "
            f"{e:.4f} | {r['auc_str']} | {d_full} | {d_ctxo} | "
            f"{fmt_p(p_full)} {m_full} | {fmt_p(p_ctxo)} {m_ctxo} |"
        )
    lines.append("")
    lines.append("Significance markers: **** p<1e-4, *** p<1e-3, ** p<1e-2, * p<5e-2, n.s. otherwise.\n")

    # Conclusions
    full_e = float(view[view["name"] == "full"]["auc_ensemble"].iloc[0])
    ctxo_e = float(view[view["name"] == "context_only"]["auc_ensemble"].iloc[0])
    seq_e = float(view[view["name"] == "sequence_only"]["auc_ensemble"].iloc[0])
    lines.append("## Conclusions\n")
    lines.append(f"- Full CARCM (all 4 features + ctx) AUC = {full_e:.4f}; "
                 f"context_only (no sequence) = {ctxo_e:.4f}; "
                 f"sequence_only (no context) = {seq_e:.4f}.")
    for n in ["only_binned", "only_positional", "only_count", "only_cooc"]:
        rr = view[view["name"] == n]
        if rr.empty:
            continue
        rr = rr.iloc[0]
        e = float(rr["auc_ensemble"])
        d_full = rr["delta_auc_vs_full_mean"]
        d_ctxo = rr["delta_auc_vs_context_only_mean"]
        lines.append(
            f"- **{rr['label']}**: ensemble AUC = {e:.4f}. "
            f"ΔAUC vs Full = {d_full:+.4f}, ΔAUC vs context_only = {d_ctxo:+.4f}. "
            f"This single feature alone, when paired with full context, recovers "
            f"{(e-ctxo_e)/(full_e-ctxo_e)*100:.1f}% of Full's gain over context_only."
        )

    lines.append("")
    lines.append("## Figures\n")
    lines.append("![LOI ΔAUC](../figures/loi_delta_auc.png)\n")
    lines.append("*Figure 1. Two-panel comparison. Left: ΔAUC vs Full CARCM (negative — "
                 "how much each single feature loses vs the all-features baseline). "
                 "Right: ΔAUC vs context_only (positive — how much each single feature "
                 "adds on top of pure context features). Bars are paired-seed mean ± 95% CI; "
                 "markers are DeLong p vs the corresponding reference.*\n")
    lines.append("![LOI metric grid](../figures/loi_metric_grid.png)\n")
    lines.append("*Figure 2. Per-configuration mean ± 95% CI for AUC, PR-AUC, F1 and MCC, "
                 "covering the 4 single-feature configurations and the 3 reference rows.*\n")

    with open(os.path.join(rep_dir, "loi_report.md"), "w") as f:
        f.write("\n".join(lines))

    try:
        write_docx_report(view, rep_dir, fig_dir, n_seeds, n_test, full_e, ctxo_e, seq_e)
    except Exception as e:
        print(f"WARN: DOCX render failed: {e}", flush=True)


def write_docx_report(view, rep_dir, fig_dir, n_seeds, n_test, full_e, ctxo_e, seq_e):
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles["Normal"]; style.font.name = "Calibri"; style.font.size = Pt(11)
    title = doc.add_heading("LOI Feature Ablation Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        f"Four leave-one-IN configurations: each retains a single sequence-side "
        f"feature block plus the full context branch. "
        f"References (Full CARCM, sequence_only, context_only) are reused from the "
        f"parent context-ablation study without retraining. Each configuration "
        f"trained with {n_seeds} seeds × 150 epochs × patience 25. Test set "
        f"n = {n_test}."
    )

    doc.add_heading("Table — Single-feature + context (LOI)", level=1)
    headers = ["#", "Configuration", "Ensemble AUC", "AUC mean [95% CI]",
               "ΔAUC vs Full", "ΔAUC vs context_only",
               "DeLong p vs Full", "DeLong p vs context_only"]
    table = doc.add_table(rows=1 + len(view), cols=len(headers))
    table.style = "Light Grid Accent 1"
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]; cell.text = ""
        run = cell.paragraphs[0].add_run(h); run.bold = True; run.font.size = Pt(10)
    for i, (_, r) in enumerate(view.iterrows(), start=1):
        e = r.get("auc_ensemble", float("nan"))
        vals = [
            str(int(r["id"])), r["label"],
            f"{e:.4f}", r["auc_str"],
            r.get("delta_auc_vs_full_str", ""),
            r.get("delta_auc_vs_context_only_str", ""),
            f"{fmt_p(r.get('delong_p_vs_full', float('nan')))} {r.get('delong_marker_vs_full','')}",
            f"{fmt_p(r.get('delong_p_vs_context_only', float('nan')))} {r.get('delong_marker_vs_context_only','')}",
        ]
        for j, v in enumerate(vals):
            cell = table.rows[i].cells[j]; cell.text = ""
            run = cell.paragraphs[0].add_run(v); run.font.size = Pt(9)

    doc.add_heading("Figures", level=1)
    f1 = os.path.join(fig_dir, "loi_delta_auc.png")
    f2 = os.path.join(fig_dir, "loi_metric_grid.png")
    if os.path.exists(f1):
        doc.add_picture(f1, width=Inches(6.5))
    if os.path.exists(f2):
        doc.add_picture(f2, width=Inches(6.5))

    doc.save(os.path.join(rep_dir, "loi_report.docx"))


if __name__ == "__main__":
    aggregate(parse_args())
